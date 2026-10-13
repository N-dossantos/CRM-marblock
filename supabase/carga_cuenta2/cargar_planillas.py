#!/usr/bin/env python3
"""
cargar_planillas.py — carga en Supabase las cuentas corrientes de Cuenta 2 que hoy viven en planillas
(Google Sheets bajadas como .xlsx a supabase/cuenta2/, carpeta gitignoreada: son datos reales).

Cuenta 2 no viene de Tango: estos datos se cargan cuando quieras y el cutover de Tango los conserva
(05_cargar.sh no trunca las tablas de Cuenta 2 — ver el comentario de TABLAS ahí).

Cada hoja es una entidad (cliente o proveedor, según el nombre del archivo: si dice "proveedor" es
compra, si no, venta). Dos formatos:

  * Cuenta corriente (fecha | remito | debo | pagos | saldo):
      - remito con importe en "debo"      -> Remito X con un renglón "importe según planilla"
      - sólo letras en "debo" ("Torres")  -> Remito X "S/N" a esa obra; un pago en la misma fila
                                             es un COBRO
      - "pg" / "pago" / "pg transf." ...  -> COBRO (venta) o PAGO (compra); medio transferencia o
                                             cheque si el texto lo dice, si no efectivo
      - "acopio lista N" en "pagos"       -> COBRO (plata adelantada)
      - "materiales" en "pagos"           -> AJUSTE al haber ("Pago con materiales")
      - hormigón / piedra / arena / m3    -> AJUSTE al haber ("Pago con ..."): pago en especie
      - pallets / publicidad / arreglo    -> AJUSTE al haber con el texto de la planilla
      - "iva fac N" / "cheque N" ...      -> AJUSTE al debe (un cheque en "debo" es un rechazado)
      - "factura N ..." / "boleta ..."    -> AJUSTE al debe con el texto de la planilla
      - saldo escrito a mano sin importes -> AJUSTE de saldo inicial (sin fecha: el día anterior al
                                             primer movimiento). Una fila "saldo" con importe en
                                             "debo" también: manda la columna saldo (se avisa si
                                             no coinciden)
      - saldo a mano en el medio, con una leyenda ("Difer pallets") -> AJUSTE por la diferencia
      - movimiento sin fecha              -> la fecha de la fila anterior (se avisa)
    Las hojas "Hoja2", "Hoja5"... sin datos se saltean. Si hay dos columnas "fecha" (otra tabla a la
    derecha, como en "Armando"), vale la primera.
  * Cantidades por producto (fecha | remito | <producto> ... , con una fila de precios abajo, como
    "Metal grande"): un Remito X por fila con sus productos del catálogo (PRODUCTOS_POR_COLUMNA).

Fechas posteriores a hoy se toman como error de tipeo del año y se cargan un año antes (se avisa).
Cada entidad se controla dos veces: el saldo que arma el script tiene que dar el de la planilla, y
la base, después de insertar, tiene que dar ese mismo saldo (si no, RAISE y vuelve todo atrás).

Una entidad que ya existe en la base CON movimientos se saltea (no se duplica). --reemplazar borra
sus remitos y movimientos de Cuenta 2 y la vuelve a cargar desde la planilla — incluidos los que se
hayan cargado a mano desde la app, así que usarlo sólo mientras la planilla siga siendo la fuente.

Uso (desde cualquier carpeta):
  python3 cargar_planillas.py                    lee todas las .xlsx de supabase/cuenta2/ y muestra
                                                 el resumen; no toca la base
  python3 cargar_planillas.py archivo.xlsx ...   idem, sólo esos archivos
  python3 cargar_planillas.py --probar           corre la carga en la base y hace ROLLBACK
  python3 cargar_planillas.py --cargar           corre la carga y hace COMMIT
  python3 cargar_planillas.py --cargar --reemplazar

La conexión sale de SUPABASE_DB_URL o, si no está, del .db_url gitignoreado de la raíz del repo.
El SQL generado queda en supabase/cuenta2/_carga.sql (gitignoreado).
"""
import argparse
import datetime as dt
import os
import re
import subprocess
import sys
from decimal import Decimal
from pathlib import Path

import openpyxl

AQUI = Path(__file__).resolve().parent
REPO = AQUI.parent.parent
CARPETA = REPO / "supabase" / "cuenta2"
SQL_SALIDA = CARPETA / "_carga.sql"
HOY = dt.date.today()
CERO = Decimal("0")

# Columnas de la planilla por producto → (productos.id, descripción del renglón). Mapeo verificado
# contra precio_sin_iva: los precios de la planilla guardan la misma proporción con la lista actual.
# "SPLIT ESQ./ DINTEL" sirve para los dos (id 16 y 17 tienen el mismo precio): va al 16 y la
# descripción conserva las dos.
PRODUCTOS_POR_COLUMNA = {
    "B20PORT": (1, "Bloque Liso de 19x19x39 Portante"),
    "BM20": (6, "Bloque Medio de 19x19x19 Liso"),
    "BU20": (7, "Bloque Dintel de 19x19x39 Liso"),
    "B15EST": (13, "Bloque Liso de 14x19x39 Estandar"),
    "B15PORT": (10, "Bloque Liso de 14x19x39 Portante"),
    "SPLIT20": (15, "Bloque de 19x19x39 Split"),
    "SPLITESQM": (19, "Bloque Medio de 19x19x39 Split Esquinero"),
    "SPLITMEDIO": (18, "Bloque Medio de 19x19x39 Split"),
    "SPLITESQ/DINTEL": (16, "Bloque de 19x19x39 Split Esquinero / Dintel"),
    "PALLETS": (28, "Pallet de Madera Vacío"),
    "B15": (10, "Bloque Liso de 14x19x39 Portante"),  # "B15" a secas (Delfica) es el portante
    "BU15": (12, "Bloque Dintel de 14x19x39 Liso"),
    "B10": (14, "Bloque Liso de 9x19x39"),
}

# Textos de la columna "pagos" que son pago en especie (Armando paga con hormigón, Fenix con piedra
# y arena) y los que son un crédito que se carga con el texto de la planilla tal cual.
EN_ESPECIE = re.compile(r"hormig|piedra|arena|\dm3\b|\bm3\b")
CREDITO_TAL_CUAL = ("pallet", "publicidad", "arreglo")
CARGO_TAL_CUAL = ("factura", "boleta")


class ErrorPlanilla(Exception):
    pass


# ── lectura ─────────────────────────────────────────────────────────────────────

def dec(v):
    """Número de la planilla → Decimal con 2 decimales; vacío o 0 → None."""
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if isinstance(v, str):
        raise ErrorPlanilla(f"se esperaba un número y hay texto: {v!r}")
    d = Decimal(str(round(float(v), 2)))
    return d if d != 0 else None


def es_formula(v):
    return isinstance(v, str) and v.startswith("=")


def norm_col(s):
    return re.sub(r"[\s.]", "", str(s)).upper()


def frase(texto):
    """'acopio      241m3' → 'Acopio 241m3'."""
    t = re.sub(r"\s+", " ", texto.strip())
    return t[:1].upper() + t[1:]


def numero_remito(v):
    """15438.0 / 'R 15438' / '15438' → '15438'; cualquier otra cosa → None."""
    if isinstance(v, (int, float)):
        return str(int(v))
    m = re.fullmatch(r"\s*R?\s*(\d+)\s*", str(v), re.I)
    return m.group(1) if m else None


def corregir_fecha(f, avisos, donde):
    if f is None:
        return None
    if not isinstance(f, dt.datetime):
        raise ErrorPlanilla(f"{donde}: fecha inválida {f!r}")
    d = f.date()
    if d > HOY:
        nueva = d.replace(year=d.year - 1)
        avisos.append(f"{donde}: fecha {d:%d/%m/%Y} en el futuro → cargada como {nueva:%d/%m/%Y}")
        d = nueva
    return d


def encabezado(ws, fila_titulos):
    """Nombre, descuento y notas: las celdas de texto arriba de la fila de títulos."""
    nombre, descuento, notas = None, CERO, []
    for r in range(1, fila_titulos):
        for c in range(1, ws.max_column + 1):
            v = ws.cell(r, c).value
            if not isinstance(v, str) or not v.strip():
                continue
            t = v.strip()
            m = re.search(r"lista\b(.*?)-\s*(\d+(?:[.,]\d+)?)\s*%", t, re.I)
            if m:
                descuento = Decimal(m.group(2).replace(",", "."))
                if re.search(r"\d", m.group(1)):  # "lista 102 -20%": la lista de precios es dato
                    notas.append(f"Lista de precios: {t}")
            elif nombre is None:
                nombre = re.sub(r"^\s*(cliente|proveedor)\s*:\s*", "", t, flags=re.I).strip()
            else:
                notas.append(f"Contacto: {t}")
    return nombre, descuento, notas


def buscar_titulos(ws):
    for r in range(1, ws.max_row + 1):
        fila = {c: str(ws.cell(r, c).value).strip().lower()
                for c in range(1, ws.max_column + 1) if ws.cell(r, c).value is not None}
        if "fecha" in fila.values():
            return r, fila
    raise ErrorPlanilla("no encuentro la fila de títulos (una celda que diga 'fecha')")


def leer_hoja(ws_f, ws_v, archivo):
    """Devuelve la entidad con sus remitos y movimientos (None si es una "HojaN" vacía), o levanta
    ErrorPlanilla."""
    hoja = f"{archivo} / {ws_f.title}"
    fila_tit, titulos = buscar_titulos(ws_f)
    # La columna saldo de una hoja vacía puede tener la fórmula copiada hasta abajo: no cuenta.
    if re.fullmatch(r"hoja\s*\d+", ws_f.title.strip(), re.I) and all(
            ws_f.cell(r, c).value is None for r in range(fila_tit + 1, ws_f.max_row + 1)
            for c, v in titulos.items() if v != "saldo"):
        return None
    nombre, descuento, notas = encabezado(ws_f, fila_tit)
    ent = {"hoja": hoja, "nombre": nombre or ws_f.title.strip(), "descuento": descuento,
           "notas": "\n".join(notas) or None, "remitos": [], "movs": [], "avisos": []}
    col = {}
    for c, v in titulos.items():  # en orden de columna: si un título se repite, vale el primero
        col.setdefault(v, c)
    if "saldo" in col:
        leer_cta_cte(ws_f, ws_v, fila_tit, col, ent)
    else:
        leer_por_producto(ws_f, ws_v, fila_tit, titulos, ent)
    return ent


def leer_cta_cte(ws_f, ws_v, fila_tit, col, ent):
    c_fecha, c_ref, c_saldo = col["fecha"], col["remito"], col["saldo"]
    c_debe = col.get("debo", col.get("debe"))
    c_haber = col.get("pagos", col.get("pago", col.get("haber")))
    if c_debe is None or c_haber is None:
        raise ErrorPlanilla(f"{ent['hoja']}: faltan las columnas debo/pagos")

    filas = []          # (fila, fecha, debe, haber, saldo_planilla) en orden de la hoja
    inicial = None      # (fila, fecha, monto) del saldo inicial escrito a mano
    ultima_fecha = None
    obras = set()
    for r in range(fila_tit + 1, ws_f.max_row + 1):
        donde = f"{ent['hoja']} fila {r}"
        saldo_f, saldo_v = ws_f.cell(r, c_saldo).value, ws_v.cell(r, c_saldo).value
        a_mano = not es_formula(saldo_f) and isinstance(saldo_f, (int, float)) and saldo_f != 0
        # Una leyenda en debo/pagos al lado de un saldo a mano ("Difer pallets") explica el ajuste.
        leyendas = [v.strip() for v in (ws_f.cell(r, c_debe).value, ws_f.cell(r, c_haber).value)
                    if isinstance(v, str) and v.strip() and not es_formula(v)]
        try:
            fecha = corregir_fecha(ws_f.cell(r, c_fecha).value, ent["avisos"], donde)
            ref = ws_f.cell(r, c_ref).value
            if leyendas and a_mano:
                debe = haber = None
            else:
                debe, haber = dec(ws_f.cell(r, c_debe).value), dec(ws_f.cell(r, c_haber).value)
        except ErrorPlanilla as e:
            raise ErrorPlanilla(f"{donde}: {e}")
        texto = str(ref).strip() if ref is not None else ""
        low = texto.lower()

        # Fila "saldo" con importe en debo: la planilla encadena desde la columna saldo, que manda.
        if low == "saldo" and a_mano and haber is None:
            if inicial is not None or filas:
                raise ErrorPlanilla(f"{donde}: fila 'saldo' en el medio de la cuenta")
            monto = Decimal(str(round(saldo_f, 2)))
            if debe is not None and debe != monto:
                ent["avisos"].append(f"{donde}: fila 'saldo' con {debe} en debo y {monto} en saldo → "
                                     f"se toma {monto}")
            inicial = (r, fecha, monto)
            continue
        if debe is None and haber is None:
            if a_mano and (inicial is not None or filas):
                if not leyendas:
                    raise ErrorPlanilla(f"{donde}: saldo escrito a mano en el medio de la cuenta")
                previo = (inicial[2] if inicial else CERO) + sum(d - h for _, d, h, _ in filas)
                dif = Decimal(str(round(saldo_f, 2))) - previo
                fecha = fecha or ultima_fecha
                ent["movs"].append(ajuste(fecha, max(dif, CERO), max(-dif, CERO), frase(leyendas[0])))
                ent["avisos"].append(f"{donde}: saldo a mano ({saldo_f}) con '{leyendas[0]}' → AJUSTE "
                                     f"de {dif} al {fecha:%d/%m/%Y}")
                filas.append((r, max(dif, CERO), max(-dif, CERO), saldo_v))
            elif a_mano:
                inicial = (r, fecha, Decimal(str(round(saldo_f, 2))))
            elif texto and low != "saldo":
                raise ErrorPlanilla(f"{donde}: '{texto}' sin importe")
            continue
        if fecha is None:
            if ultima_fecha is None:
                raise ErrorPlanilla(f"{donde}: movimiento sin fecha")
            fecha = ultima_fecha
            ent["avisos"].append(f"{donde}: '{texto}' sin fecha → {fecha:%d/%m/%Y} (la de la fila anterior)")
        ultima_fecha = fecha

        nro = numero_remito(ref) if ref is not None else None
        obra = None
        if debe:
            if nro:
                ent["remitos"].append({"numero": nro, "fecha": fecha, "total": debe, "items": None})
            elif low.startswith("iva fac"):
                m = re.search(r"\d+", texto)
                concepto = f"IVA factura {m.group(0)}" if m else texto
                ent["movs"].append(ajuste(fecha, debe, CERO, concepto))
            elif "cheque" in low:  # un cheque que suma al debe es uno rechazado
                m = re.search(r"\d+", texto)
                concepto = "Cheque rechazado" + (f" {m.group(0)}" if m else "")
                ent["movs"].append(ajuste(fecha, debe, CERO, concepto))
            elif low.startswith(CARGO_TAL_CUAL):
                ent["movs"].append(ajuste(fecha, debe, CERO, frase(texto)))
            elif re.fullmatch(r"[a-záéíóúüñ ]+", low):  # sin número: un remito a la obra "Torres"
                obra = " ".join(texto.split()).title()
                obras.add(obra)
                ent["remitos"].append({"numero": "S/N", "fecha": fecha, "total": debe, "items": None,
                                       "detalle": f"Obra {obra}: importe según planilla (sin detalle de productos)"})
            else:
                raise ErrorPlanilla(f"{donde}: no sé qué es '{texto}' en 'debo'")
        if haber:
            if "material" in low:
                m = re.search(r"\d+", texto)
                concepto = "Pago con materiales" + (f" ({m.group(0)})" if m else "")
                ent["movs"].append(ajuste(fecha, CERO, haber, concepto))
            elif EN_ESPECIE.search(low):
                concepto = frase(texto) if low.startswith("acopio") else "Pago con " + " ".join(low.split())
                ent["movs"].append(ajuste(fecha, CERO, haber, concepto))
            elif any(p in low for p in CREDITO_TAL_CUAL):
                ent["movs"].append(ajuste(fecha, CERO, haber, frase(texto)))
            elif low.startswith("acopio"):
                ent["movs"].append(cobro(fecha, haber, frase(texto)))
            elif nro or obra or low in ("pg", "pago") or re.match(r"(pg|pago)\b", low):
                ent["movs"].append(cobro(fecha, haber, texto if not (nro or obra) else "pago"))
            else:
                raise ErrorPlanilla(f"{donde}: no sé qué es '{texto}' en 'pagos'")
        filas.append((r, debe or CERO, haber or CERO, saldo_v))
    if obras:
        n = sum(1 for x in ent["remitos"] if x["numero"] == "S/N")
        ent["avisos"].append(f"{ent['hoja']}: {n} remitos sin número, a las obras {', '.join(sorted(obras))}")

    # Saldo inicial sin fecha: el día anterior al primer movimiento.
    if inicial:
        r, fecha, monto = inicial
        if fecha is None:
            primera = min([x["fecha"] for x in ent["remitos"]] + [x["fecha"] for x in ent["movs"]])
            fecha = primera - dt.timedelta(days=1)
            ent["avisos"].append(f"{ent['hoja']} fila {r}: saldo inicial sin fecha → {fecha:%d/%m/%Y}")
        mov = ajuste(fecha, monto if monto > 0 else CERO, -monto if monto < 0 else CERO,
                     "Saldo inicial según planilla")
        ent["movs"].insert(0, mov)

    # Control 1: la suma tiene que reproducir la columna saldo de la planilla, fila por fila.
    saldo = inicial[2] if inicial else CERO
    for r, debe, haber, saldo_v in filas:
        saldo += debe - haber
        if isinstance(saldo_v, (int, float)) and abs(Decimal(str(round(saldo_v, 2))) - saldo) > Decimal("0.01"):
            raise ErrorPlanilla(f"{ent['hoja']} fila {r}: el saldo calculado ({saldo}) no da el de la "
                                f"planilla ({saldo_v})")
    ent["saldo"] = saldo


def ajuste(fecha, debe, haber, concepto):
    return {"fecha": fecha, "tipo": "AJUSTE", "medio": None, "debe": debe, "haber": haber,
            "concepto": concepto}


def cobro(fecha, monto, texto):
    low = texto.lower()
    if "transf" in low:
        medio = "transferencia"
    elif "cheque" in low and "rech" not in low:
        medio = "cheque"  # el cheque en sí no está en la planilla: el movimiento queda sin cheque_id
    else:
        medio = "efectivo"
    if low in ("pg", "pago"):
        concepto = "Pago"
    else:
        concepto = re.sub(r"^pg\b\.?", "Pago", texto.strip(), flags=re.I)
        concepto = concepto[:1].upper() + concepto[1:]
    return {"fecha": fecha, "tipo": "COBRO", "medio": medio, "debe": CERO, "haber": monto,
            "concepto": concepto}


def leer_por_producto(ws_f, ws_v, fila_tit, titulos, ent):
    c_fecha = next(c for c, v in titulos.items() if v == "fecha")
    c_rem = next(c for c, v in titulos.items() if v == "remito")
    cols = {}
    for c, v in titulos.items():
        if c in (c_fecha, c_rem):
            continue
        clave = norm_col(ws_f.cell(fila_tit, c).value)
        if clave not in PRODUCTOS_POR_COLUMNA:
            raise ErrorPlanilla(f"{ent['hoja']}: columna de producto desconocida "
                                f"'{ws_f.cell(fila_tit, c).value}' (agregala a PRODUCTOS_POR_COLUMNA)")
        cols[c] = PRODUCTOS_POR_COLUMNA[clave]

    filas, precios, r = [], None, fila_tit + 1
    while r <= ws_f.max_row:
        f = ws_f.cell(r, c_fecha).value
        if isinstance(f, dt.datetime):
            filas.append(r)
        elif filas and precios is None:
            vals = [ws_f.cell(r, c).value for c in cols]
            if any(isinstance(v, (int, float)) for v in vals) and not any(es_formula(v) for v in vals):
                precios = r
        r += 1
    if precios is None:
        raise ErrorPlanilla(f"{ent['hoja']}: no encuentro la fila de precios debajo de los remitos")

    total = CERO
    for r in filas:
        donde = f"{ent['hoja']} fila {r}"
        nro = numero_remito(ws_f.cell(r, c_rem).value)
        if not nro:
            raise ErrorPlanilla(f"{donde}: remito sin número")
        items = []
        for c, (prod_id, desc) in cols.items():
            cant = dec(ws_f.cell(r, c).value)
            if cant:
                precio = dec(ws_f.cell(precios, c).value)
                if precio is None:
                    raise ErrorPlanilla(f"{donde}: '{desc}' sin precio en la fila {precios}")
                items.append({"producto_id": prod_id, "descripcion": desc, "cantidad": cant,
                              "precio": precio, "subtotal": (cant * precio).quantize(Decimal("0.01"))})
        if not items:
            continue
        sub = sum(i["subtotal"] for i in items)
        total += sub
        fecha = corregir_fecha(ws_f.cell(r, c_fecha).value, ent["avisos"], donde)
        ent["remitos"].append({"numero": nro, "fecha": fecha, "total": sub, "items": items})

    # Control 1: el total general de la planilla (la última SUMA de rango, debajo de los precios).
    esperado = None
    for r in range(precios + 1, ws_f.max_row + 1):
        for c in range(1, ws_f.max_column + 1):
            f, v = ws_f.cell(r, c).value, ws_v.cell(r, c).value
            if es_formula(f) and ":" in f and isinstance(v, (int, float)):
                esperado = Decimal(str(round(v, 2)))
    if esperado is not None and esperado != total:
        raise ErrorPlanilla(f"{ent['hoja']}: los remitos suman {total} y la planilla dice {esperado}")
    ent["saldo"] = total


# ── SQL ─────────────────────────────────────────────────────────────────────────

def q(v):
    if v is None:
        return "NULL"
    if isinstance(v, bool):
        return "true" if v else "false"
    if isinstance(v, (Decimal, int)):
        return str(v)
    if isinstance(v, dt.date):
        return f"'{v.isoformat()}'"
    return "'" + str(v).replace("'", "''") + "'"


def sql_entidad(ent, sector, reemplazar):
    tabla = "clientes_cuenta2" if sector == "venta" else "proveedores_cuenta2"
    fk = "cliente_id" if sector == "venta" else "proveedor_id"
    nombre = q(ent["nombre"])
    obs = q(f"Importado de planilla: {ent['hoja']}")
    L = [f"-- {ent['hoja']}",
         "DO $do$",
         "DECLARE v_ent integer; v_rem integer; v_n integer; v_saldo numeric;",
         "BEGIN",
         f"  SELECT id INTO v_ent FROM {tabla} WHERE lower(btrim(nombre)) = lower(btrim({nombre}));",
         "  IF v_ent IS NOT NULL THEN",
         f"    SELECT (SELECT count(*) FROM cuenta2_remitos WHERE {fk} = v_ent)",
         f"         + (SELECT count(*) FROM cuenta2_movimientos WHERE {fk} = v_ent) INTO v_n;",
         "    IF v_n > 0 THEN",
         f"      IF NOT {q(reemplazar)} THEN",
         f"        RAISE NOTICE 'SALTEADO %: ya tiene % remitos/movimientos (usá --reemplazar)', {nombre}, v_n;",
         "        RETURN;",
         "      END IF;",
         # Un cheque de la cartera C2 ligado a la entidad quedaría huérfano: mejor frenar.
         f"      IF EXISTS (SELECT 1 FROM cuenta2_cheques WHERE {fk} = v_ent OR origen_movimiento_id IN",
         f"                 (SELECT id FROM cuenta2_movimientos WHERE {fk} = v_ent)) THEN",
         f"        RAISE EXCEPTION '%: tiene cheques en la cartera de Cuenta 2, no se reemplaza a ciegas', {nombre};",
         "      END IF;",
         f"      DELETE FROM cuenta2_movimientos WHERE {fk} = v_ent;",
         f"      DELETE FROM cuenta2_remitos WHERE {fk} = v_ent;",
         "    END IF;",
         f"    UPDATE {tabla} SET descuento_porcentaje = {q(ent['descuento'])},",
         f"                       notas = COALESCE(notas, {q(ent['notas'])}) WHERE id = v_ent;",
         "  ELSE",
         f"    INSERT INTO {tabla} (nombre, descuento_porcentaje, notas)",
         f"    VALUES ({nombre}, {q(ent['descuento'])}, {q(ent['notas'])}) RETURNING id INTO v_ent;",
         "  END IF;"]
    for rem in ent["remitos"]:
        L.append("  INSERT INTO cuenta2_remitos (tipo_sector, " + fk + ", numero, fecha, descuento_porcentaje,"
                 " subtotal, descuento_monto, total, observaciones)")
        L.append(f"  VALUES ({q(sector)}, v_ent, {q(rem['numero'])}, {q(rem['fecha'])}, 0, {q(rem['total'])}, 0,"
                 f" {q(rem['total'])}, {obs}) RETURNING id INTO v_rem;")
        items = rem["items"] or [{"producto_id": None,
                                  "descripcion": rem.get("detalle", "Importe según planilla (sin detalle de productos)"),
                                  "cantidad": Decimal("1"), "precio": rem["total"], "subtotal": rem["total"]}]
        valores = ",\n    ".join(
            f"(v_rem, {q(i['producto_id'])}, {q(i['descripcion'])}, {q(i['cantidad'])}, {q(i['precio'])}, 0,"
            f" {q(i['subtotal'])}, {n})" for n, i in enumerate(items))
        L.append("  INSERT INTO cuenta2_remito_items (remito_id, producto_id, descripcion, cantidad,"
                 " precio_unitario, descuento_item, subtotal, orden) VALUES")
        L.append(f"    {valores};")
    if ent["movs"]:
        tipo_cobro = "COBRO" if sector == "venta" else "PAGO"
        valores = ",\n    ".join(
            f"({q(sector)}, v_ent, {q(m['fecha'])}, {q(tipo_cobro if m['tipo'] == 'COBRO' else m['tipo'])},"
            f" {q(m['medio'])}, {q(m['debe'])}, {q(m['haber'])}, {q(m['concepto'])})" for m in ent["movs"])
        L.append(f"  INSERT INTO cuenta2_movimientos (tipo_sector, {fk}, fecha, tipo, medio, debe, haber, concepto) VALUES")
        L.append(f"    {valores};")
    # Control 2: el saldo según la base (misma cuenta que clientes_cuenta2_list).
    L += [f"  SELECT COALESCE((SELECT SUM(total) FROM cuenta2_remitos WHERE {fk} = v_ent), 0)",
          f"       + COALESCE((SELECT SUM(debe - haber) FROM cuenta2_movimientos WHERE {fk} = v_ent), 0) INTO v_saldo;",
          f"  IF v_saldo <> {q(ent['saldo'])} THEN",
          f"    RAISE EXCEPTION '%: la base da saldo % y la planilla {ent['saldo']}', {nombre}, v_saldo;",
          "  END IF;",
          f"  RAISE NOTICE 'OK %: {len(ent['remitos'])} remitos, {len(ent['movs'])} movimientos, saldo %', {nombre}, v_saldo;",
          "END $do$;", ""]
    return "\n".join(L)


# ── main ────────────────────────────────────────────────────────────────────────

def db_url():
    url = os.environ.get("SUPABASE_DB_URL")
    if not url and (REPO / ".db_url").exists():
        url = (REPO / ".db_url").read_text().strip()
    if not url:
        sys.exit("ERROR: falta SUPABASE_DB_URL (o el archivo .db_url en la raíz del repo).")
    return url


def main():
    sys.stdout.reconfigure(line_buffering=True)  # que el resumen salga antes que los NOTICE de psql
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("archivos", nargs="*", type=Path)
    modo = ap.add_mutually_exclusive_group()
    modo.add_argument("--probar", action="store_true", help="corre en la base y hace ROLLBACK")
    modo.add_argument("--cargar", action="store_true", help="corre en la base y hace COMMIT")
    ap.add_argument("--reemplazar", action="store_true",
                    help="borra y recarga las entidades que ya tienen datos en Cuenta 2")
    args = ap.parse_args()

    archivos = args.archivos or sorted(p for p in CARPETA.glob("*.xlsx") if not p.name.startswith("~$"))
    if not archivos:
        sys.exit(f"No hay .xlsx en {CARPETA}")

    entidades, errores = [], []
    for arch in archivos:
        sector = "compra" if "proveedor" in arch.name.lower() else "venta"
        wb_f = openpyxl.load_workbook(arch, data_only=False)
        wb_v = openpyxl.load_workbook(arch, data_only=True)
        for ws in wb_f.worksheets:
            try:
                ent = leer_hoja(ws, wb_v[ws.title], arch.name)
                if ent is None:
                    continue
                ent["sector"] = sector
                entidades.append(ent)
            except ErrorPlanilla as e:
                errores.append(str(e))

    vistos = {}
    for e in entidades:
        clave = (e["sector"], e["nombre"].strip().lower())
        if clave in vistos:
            errores.append(f"'{e['nombre']}' aparece dos veces: {vistos[clave]} y {e['hoja']}")
        vistos[clave] = e["hoja"]

    print(f"{'Entidad':<40} {'Sector':<7} {'Dto':>5} {'Remitos':>8} {'Movs':>5} {'Saldo':>18}")
    for e in entidades:
        print(f"{e['nombre']:<40} {e['sector']:<7} {e['descuento']:>4}% {len(e['remitos']):>8} "
              f"{len(e['movs']):>5} {e['saldo']:>18,.2f}")
    avisos = [a for e in entidades for a in e["avisos"]]
    if avisos:
        print("\nAvisos:\n  " + "\n  ".join(avisos))
    if errores:
        print("\nERRORES (no se genera nada hasta corregirlos):\n  " + "\n  ".join(errores))
        sys.exit(1)

    fin = "COMMIT;" if args.cargar else "ROLLBACK;"
    sql = ["-- Generado por supabase/carga_cuenta2/cargar_planillas.py. Datos reales: no versionar.",
           "BEGIN;", ""] + [sql_entidad(e, e["sector"], args.reemplazar) for e in entidades] + [fin, ""]
    SQL_SALIDA.write_text("\n".join(sql), encoding="utf-8")
    print(f"\nSQL: {SQL_SALIDA.relative_to(REPO)}")

    if not (args.probar or args.cargar):
        print("No se tocó la base. --probar para ensayarla con ROLLBACK, --cargar para grabarla.")
        return
    r = subprocess.run(["psql", db_url(), "-v", "ON_ERROR_STOP=1", "-X", "-q", "-f", str(SQL_SALIDA)])
    if r.returncode != 0:
        sys.exit("\nFalló: la transacción volvió atrás, la base quedó como estaba.")
    print("\nCargado (COMMIT)." if args.cargar else "\nEnsayo OK — ROLLBACK, la base quedó como estaba.")


if __name__ == "__main__":
    main()
